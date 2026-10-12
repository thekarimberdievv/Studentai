import os, io, hmac, hashlib, json, time, math, asyncio
from urllib.parse import parse_qsl
from contextlib import asynccontextmanager
import asyncpg, httpx
from docx import Document
from fastapi import FastAPI, Header, HTTPException, Request
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

BOT_TOKEN = os.environ["BOT_TOKEN"]
DB_URL = os.environ["DATABASE_URL"]
BASE_URL = os.getenv("BASE_URL", "").rstrip("/")
ADMINS = {int(x) for x in os.getenv("ADMIN_IDS", "").split(",") if x.strip()}
AI_KEY = os.getenv("AI_API_KEY", "")
AI_MODEL = os.getenv("AI_MODEL", "claude-sonnet-4-6")
WH_SECRET = os.getenv("WEBHOOK_SECRET", "change-me")
DEV = os.getenv("DEV_MODE") == "1"
TG = f"https://api.telegram.org/bot{BOT_TOKEN}"

SERVICES = {
    "referat": ("Referat", "bet", 3000), "mustaqil": ("Mustaqil ish", "bet", 3000),
    "kurs": ("Kurs ishi", "bet", 8000), "slayd": ("Slayd / taqdimot", "slayd", 2000),
    "esse": ("Esse", "bet", 3000), "maqola": ("Ilmiy maqola", "bet", 6000),
    "konspekt": ("Konspekt", "bet", 2000), "test": ("Test tuzish", "savol", 500),
    "amaliy": ("Amaliy ish / hisobot", "bet", 4000), "diplom": ("Diplom ishi", "bet", 15000),
}
CX = {1.0, 1.3, 1.7}
pool: asyncpg.Pool
http: httpx.AsyncClient
tasks: set = set()


@asynccontextmanager
async def lifespan(app):
    global pool, http
    pool = await asyncpg.create_pool(DB_URL)
    http = httpx.AsyncClient(timeout=180)
    async with pool.acquire() as c:
        await c.execute("""
        CREATE TABLE IF NOT EXISTS users(tg_id BIGINT PRIMARY KEY, name TEXT, balance BIGINT DEFAULT 0, created TIMESTAMPTZ DEFAULT now());
        CREATE TABLE IF NOT EXISTS prices(svc TEXT PRIMARY KEY, price BIGINT);
        CREATE TABLE IF NOT EXISTS orders(id SERIAL PRIMARY KEY, tg_id BIGINT, svc TEXT, topic TEXT, qty INT, cx REAL,
            cost BIGINT, status INT DEFAULT 0, created TIMESTAMPTZ DEFAULT now());""")
        for r in await c.fetch("SELECT id FROM orders WHERE status IN (0,1)"):
            spawn(process(r["id"]))
    if BASE_URL:
        await http.post(f"{TG}/setWebhook", json={"url": f"{BASE_URL}/webhook/{WH_SECRET}"})
        await http.post(f"{TG}/setChatMenuButton", json={"menu_button": {"type": "web_app", "text": "Student AI", "web_app": {"url": BASE_URL}}})
    yield
    await http.aclose(); await pool.close()


app = FastAPI(lifespan=lifespan)


def spawn(coro):
    t = asyncio.create_task(coro); tasks.add(t); t.add_done_callback(tasks.discard)


def verify(init_data: str) -> dict:
    try:
        d = dict(parse_qsl(init_data, keep_blank_values=True))
        h = d.pop("hash")
        s = "\n".join(f"{k}={v}" for k, v in sorted(d.items()))
        key = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
        if not hmac.compare_digest(hmac.new(key, s.encode(), hashlib.sha256).hexdigest(), h):
            raise ValueError
        if time.time() - int(d["auth_date"]) > 86400:
            raise ValueError
        return json.loads(d["user"])
    except Exception:
        raise HTTPException(401, "Telegram orqali oching")


async def auth(init_data: str) -> dict:
    u = verify(init_data)
    name = (u.get("first_name") or "Talaba")[:60]
    async with pool.acquire() as c:
        await c.execute("INSERT INTO users(tg_id,name) VALUES($1,$2) ON CONFLICT(tg_id) DO UPDATE SET name=$2", u["id"], name)
    return {"id": u["id"], "name": name}


async def prices():
    async with pool.acquire() as c:
        over = {r["svc"]: r["price"] for r in await c.fetch("SELECT * FROM prices")}
    return {k: over.get(k, v[2]) for k, v in SERVICES.items()}


@app.get("/api/me")
async def me(x_init_data: str = Header("")):
    u = await auth(x_init_data)
    async with pool.acquire() as c:
        bal = await c.fetchval("SELECT balance FROM users WHERE tg_id=$1", u["id"])
        rows = await c.fetch("SELECT * FROM orders WHERE tg_id=$1 ORDER BY id DESC LIMIT 50", u["id"])
    return {"name": u["name"], "bal": bal, "admin": u["id"] in ADMINS, "prices": await prices(),
            "orders": [order_out(r) for r in rows]}


def order_out(r):
    n, unit, _ = SERVICES[r["svc"]]
    return {"id": r["id"], "n": n, "t": r["topic"], "q": r["qty"], "u": unit, "c": r["cost"], "s": r["status"]}


class OrderIn(BaseModel):
    svc: str
    topic: str
    qty: int
    cx: float


@app.post("/api/order")
async def new_order(o: OrderIn, x_init_data: str = Header("")):
    u = await auth(x_init_data)
    topic = o.topic.strip()[:300]
    if o.svc not in SERVICES or not topic or not 1 <= o.qty <= 100 or o.cx not in CX:
        raise HTTPException(400, "Ma'lumot noto'g'ri")
    cost = round(o.qty * (await prices())[o.svc] * o.cx)
    async with pool.acquire() as c, c.transaction():
        bal = await c.fetchval("SELECT balance FROM users WHERE tg_id=$1 FOR UPDATE", u["id"])
        if bal < cost:
            raise HTTPException(402, "Balans yetarli emas. Balansni to'ldiring.")
        await c.execute("UPDATE users SET balance=balance-$1 WHERE tg_id=$2", cost, u["id"])
        oid = await c.fetchval("INSERT INTO orders(tg_id,svc,topic,qty,cx,cost) VALUES($1,$2,$3,$4,$5,$6) RETURNING id",
                               u["id"], o.svc, topic, o.qty, o.cx, cost)
    spawn(process(oid))
    return {"id": oid}


class TopIn(BaseModel):
    amount: int
    tg_id: int | None = None


@app.post("/api/topup")
async def topup(b: TopIn, x_init_data: str = Header("")):
    u = await auth(x_init_data)
    if not DEV:
        raise HTTPException(403, "To'lov tizimi hali ulanmagan. Balansni admin to'ldiradi.")
    await credit(u["id"], min(max(b.amount, 0), 1_000_000))
    return {"ok": 1}


async def credit(tg_id, amount):
    async with pool.acquire() as c:
        await c.execute("UPDATE users SET balance=balance+$1 WHERE tg_id=$2", amount, tg_id)


async def admin(init_data):
    u = await auth(init_data)
    if u["id"] not in ADMINS:
        raise HTTPException(403, "Ruxsat yo'q")


@app.get("/api/admin/orders")
async def a_orders(x_init_data: str = Header("")):
    await admin(x_init_data)
    async with pool.acquire() as c:
        rows = await c.fetch("SELECT o.*, u.name FROM orders o JOIN users u USING(tg_id) ORDER BY id DESC LIMIT 100")
    return [{**order_out(r), "user": r["name"]} for r in rows]


class PriceIn(BaseModel):
    svc: str
    price: int


@app.put("/api/admin/price")
async def a_price(b: PriceIn, x_init_data: str = Header("")):
    await admin(x_init_data)
    if b.svc not in SERVICES or b.price < 0:
        raise HTTPException(400, "Noto'g'ri")
    async with pool.acquire() as c:
        await c.execute("INSERT INTO prices VALUES($1,$2) ON CONFLICT(svc) DO UPDATE SET price=$2", b.svc, b.price)
    return {"ok": 1}


@app.post("/api/admin/topup")
async def a_topup(b: TopIn, x_init_data: str = Header("")):
    await admin(x_init_data)
    if not b.tg_id:
        raise HTTPException(400, "Telegram ID kerak")
    await credit(b.tg_id, b.amount)
    await tg_send(b.tg_id, f"Balansingiz {b.amount:,} so'mga to'ldirildi.".replace(",", " "))
    return {"ok": 1}


async def ai(prompt: str) -> str:
    r = await http.post("https://api.anthropic.com/v1/messages",
                        headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01"},
                        json={"model": AI_MODEL, "max_tokens": 4000, "messages": [{"role": "user", "content": prompt}]})
    r.raise_for_status()
    return "".join(b.get("text", "") for b in r.json()["content"])


async def generate(svc, topic, qty) -> str:
    name, unit, _ = SERVICES[svc]
    step = 6 if unit == "bet" else 10
    parts, tail = [], ""
    for i in range(math.ceil(qty / step)):
        n = min(step, qty - i * step)
        hint = {"bet": f"taxminan {n * 300} so'z", "slayd": f"{n} ta slayd (har biri sarlavha + 3-5 band)",
                "savol": f"{n} ta test savoli, har biri 4 variant (A-D), to'g'ri javobi bilan"}[unit]
        prompt = (f"O'zbek tilida akademik '{name}' yozing. Mavzu: {topic}.\n"
                  f"Bu {i + 1}/{math.ceil(qty / step)}-qism, hajmi: {hint}. Markdown sarlavhalari (#, ##) ishlating.\n"
                  + (f"Oldingi qism oxiri (takrorlamang, davom eting):\n{tail}\n" if tail else "Kirish qismidan boshlang.\n")
                  + ("Oxirgi qism: xulosa va foydalanilgan adabiyotlar ro'yxatini qo'shing." if i == math.ceil(qty / step) - 1 else ""))
        txt = await ai(prompt)
        parts.append(txt); tail = txt[-400:]
    return "\n\n".join(parts)


def to_docx(text: str) -> bytes:
    d = Document()
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            d.add_heading(s.lstrip("#").strip(), level=min(len(s) - len(s.lstrip("#")), 3))
        else:
            d.add_paragraph(s.replace("**", ""))
    b = io.BytesIO(); d.save(b); return b.getvalue()


async def process(oid: int):
    async with pool.acquire() as c:
        o = await c.fetchrow("UPDATE orders SET status=1 WHERE id=$1 AND status IN (0,1) RETURNING *", oid)
    if not o:
        return
    try:
        text = await generate(o["svc"], o["topic"], o["qty"])
        name = SERVICES[o["svc"]][0]
        await http.post(f"{TG}/sendDocument", data={"chat_id": o["tg_id"], "caption": f"✅ {name} tayyor: {o['topic'][:150]}"},
                        files={"document": (f"{name}-{oid}.docx", to_docx(text))})
        async with pool.acquire() as c:
            await c.execute("UPDATE orders SET status=2 WHERE id=$1", oid)
    except Exception as e:
        print("order failed", oid, repr(e))
        async with pool.acquire() as c, c.transaction():
            await c.execute("UPDATE orders SET status=3 WHERE id=$1", oid)
            await c.execute("UPDATE users SET balance=balance+$1 WHERE tg_id=$2", o["cost"], o["tg_id"])
        await tg_send(o["tg_id"], "Buyurtmani bajarib bo'lmadi. Pul balansingizga qaytarildi, qayta urinib ko'ring.")


async def tg_send(chat_id, text, markup=None):
    p = {"chat_id": chat_id, "text": text}
    if markup:
        p["reply_markup"] = markup
    await http.post(f"{TG}/sendMessage", json=p)


@app.post("/webhook/{secret}")
async def webhook(secret: str, req: Request):
    if secret != WH_SECRET:
        raise HTTPException(404)
    m = (await req.json()).get("message") or {}
    if m.get("chat", {}).get("type") == "private":
        await tg_send(m["chat"]["id"],
                      "🎓 Student AI — referat, slayd, esse, test va boshqa akademik ishlar sun'iy intellekt yordamida.\nBoshlash uchun pastdagi tugmani bosing.",
                      {"inline_keyboard": [[{"text": "Ochish", "web_app": {"url": BASE_URL}}]]})
    return {"ok": True}


app.mount("/", StaticFiles(directory="static", html=True), name="static")
